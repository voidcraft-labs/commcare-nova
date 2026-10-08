"""Defect 5: what HQ's lookup table upload answers Nova's push without the privilege, and for a 32-character tag.

Nova pushes a workbook to HQ's lookup table API before it sends an app that
reads the tables (``lib/commcare/hq/lookupTables.ts::uploadLookupTableWorkbook``).
Two things defect 5 says of that view are observed here on the view itself,
answered through HQ's URLconf, middleware and decorators with an API key, to
the bytes Nova's own client sent (the corpus's captured request).

- **Without the Lookup Tables privilege.** Contract: HQ answers the push
  with its "Upgrade Required" page, an HTML document with HTTP status 200,
  and writes no table (``accounting/decorators.py::requires_privilege_with_fallback``,
  ``domain/views/accounting.py::SubscriptionUpgradeRequiredView``), where a
  client expecting the API's JSON reads a success status. The plausible
  failures: a 402, as the decorator's own comment says it answers; a JSON
  refusal; a table written all the same. The accepted counterpart is the
  same request in the document's own project space, which HQ answers with
  the API's JSON and a table.
- **A 32-character tag.** Contract: HQ's workbook reader takes a table
  whose tag, and so whose sheet name, holds 32 characters
  (``fixtures/upload/workbook.py``), one more than a spreadsheet program
  lets a sheet name hold and one more than Nova's push allows
  (``lib/commcare/lookup/workbook.ts::MAX_HQ_FIXTURE_SHEET_NAME_LENGTH``).
  The workbook is Nova's own with its one table renamed. Beside it, a tag
  of 33 characters: HQ's upload checks no length, its column holds 32
  characters, and the view fails with a database error HQ answers as a 500,
  so the cap a client needs is 32, and it is the client's to hold.

The page's own script is named by HQ's webpack manifest, which a
deployment's static build writes and the image's holds only for the pages
its browsers run. ``base_page_script`` names the base page's entry in it
for the block, as a deployment's manifest does; the answer's status, type
and text are HQ's.
"""

from __future__ import annotations

import dataclasses
import io
import json
from contextlib import contextmanager
from unittest import mock

from proof.views.conftest import api_key, ask

DOCUMENT = "lookup-app"
PRIVILEGE = "LOOKUP_TABLES"
# The entry every HQ page's base template names (``hqwebapp/base.html``).
BASE_ENTRY = "hqwebapp/js/base"


@contextmanager
def base_page_script():
    """HQ's webpack manifests naming the base page's script, as a deployment's build writes them: the image's
    manifests hold entries only for the pages its browsers run, and a page's template refuses to render without
    its entry (``hq_shared_tags.py::webpack_bundles``)."""
    from corehq.apps.hqwebapp.utils import webpack

    read = webpack.get_webpack_manifest

    def with_base(filename=None):
        return {BASE_ENTRY: [f"{BASE_ENTRY}.js"], **read(filename)}

    with mock.patch.object(webpack, "get_webpack_manifest", with_base):
        yield


def _tags(unit):
    from corehq.apps.fixtures.models import LookupTable

    return sorted(table.tag for table in LookupTable.objects.filter(domain=unit.domain))


def _push(unit, captured, body=None, label="push"):
    """Nova's captured lookup push, sent to HQ as Nova's client sends it, with an API key of the web user."""
    return ask(
        unit,
        captured.meta["method"],
        captured.meta["path"],
        body=captured.body_path.read_bytes() if body is None else body,
        headers=(("Content-Type", captured.content_type), ("Authorization", api_key(unit))),
        label=label,
    )


def test_hq_answers_the_push_with_its_upgrade_page_and_status_200_where_the_plan_lacks_lookup_tables(
    hq, core_runner, view_documents, published
):
    document = view_documents[DOCUMENT]
    granted = document.exports["minimum"].configuration.hq()
    assert PRIVILEGE in granted.privileges
    withheld = dataclasses.replace(granted, privileges=granted.privileges - {PRIVILEGE})

    with published(document, "minimum", core_runner, create=False) as (unit, _, export):
        answer = _push(unit, export.create.lookups)
        assert (answer.status, answer.content_type) == (200, "application/json"), answer
        assert json.loads(answer.body)["code"] == 200
        written = _tags(unit)
        assert written, "HQ wrote no table for the push it accepted"

    with published(document, "minimum", core_runner, configuration=withheld, create=False) as (unit, _, export):
        with base_page_script():
            answer = _push(unit, export.create.lookups)
        assert answer.status == 200, answer.status
        assert answer.content_type.startswith("text/html"), answer.content_type
        assert b"Upgrade Required" in answer.body
        assert _tags(unit) == []


def _renamed(workbook: bytes, tag: str) -> bytes:
    """Nova's workbook with its one table's tag, and the sheet that holds its rows, renamed ``tag``."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(workbook))
    types = book["types"]
    header = [cell.value for cell in types[1]]
    column = header.index("table_id") + 1
    (old,) = {types.cell(row=row, column=column).value for row in range(2, types.max_row + 1)}
    for row in range(2, types.max_row + 1):
        types.cell(row=row, column=column).value = tag
    book[old].title = tag
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _multipart(captured, workbook: bytes) -> bytes:
    """The captured push with its workbook part replaced by ``workbook``, every other byte as Nova sent it."""
    from proof.hq.operations import upload_field

    sent = captured.body_path.read_bytes()
    held = upload_field(captured.upload(), "file-to-upload")
    assert sent.count(held) == 1
    return sent.replace(held, workbook)


def test_hqs_upload_reads_a_table_whose_tag_holds_32_characters_and_fails_on_33(
    hq, core_runner, view_documents, published
):
    document = view_documents[DOCUMENT]
    found = {}
    for length in (32, 33):
        tag = ("region_" * 5)[:length]
        assert len(tag) == length
        with published(document, "minimum", core_runner, create=False) as (unit, _, export):
            captured = export.create.lookups
            answer = _push(unit, captured, _multipart(captured, _renamed(captured.workbook(), tag)), f"tag-{length}")
            code = json.loads(answer.body)["code"] if answer.content_type == "application/json" else None
            found[length] = (answer.status, code, _tags(unit) == [tag])
    # HQ's column holds 32 characters (``LookupTable.tag``), and its upload checks no length before it writes:
    # a longer tag is a database error HQ answers with a 500, never a refusal a client can read.
    assert found == {32: (200, 200, True), 33: (500, None, False)}, found
