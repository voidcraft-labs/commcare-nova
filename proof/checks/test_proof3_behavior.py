"""Proof 3, behavioral equivalence: the two export paths, and two publishes that still build differently, behave alike.

Contract (plan work item 11, proof 3, and decision 15), judged from the
document's records (the sessions and HQ's processing are observed over B's
state, or A's where HQ refused B, ``proof.observe.sessions``): over the document's
case database, the sessions the Core runner derives on ``build(A)`` replay on
Nova's local archive, and on the build of B aligned to A wherever proof 2
still finds a difference, with identical traces after the registered spelling
rules; the profile is compared on its required version and the properties a
runtime running both builds reads; and HQ's case processing reads the same
case blocks from both sides' submissions. Sessions run only on an HQ build a
runtime can install. Every difference must fall in a class of
``proof/known-defects.json``.

The plausible failures each test here catches:

- a trace that omits a difference (the plan's row "The Core runner's traces
  are faithful"), on either route: one answer path altered in one form of a
  corpus document's build of A must give a different trace and different
  case processing, through proof 3's own comparison, in every run that
  answers the question and in no run that does not open its form. On the B
  route the altered build stands as B; proof 3 runs B's sessions only where
  proof 2 still differs, and the same build twice gives none. On the local
  route, the comparison that always runs, A's build with every form's
  ``xmlns`` replaced stands as Nova's local archive: it must compare equal
  (the namespace mapping, the local restore and the submission's namespaces
  hide nothing and invent nothing), and its altered copy must differ exactly
  where the B route's does;
- sessions run on a build HQ refuses to make (``validate_app`` lists an
  error), or a lookup table HQ cannot serve escaping as an exception the
  register can never hold; a refusal whose cause the bar already reports
  reported again, or one the bar does not report whose path names no cause,
  so two causes are one class;
- HQ's case blocks compared by their place in the form, or a case's blocks by
  their place in the case's sequence, where HQ applies them by case id
  (``casexml/apps/case/xform.py::get_case_updates``) and the form holds each
  at its own place: two blocks the builds write in another order, or one
  block more or less, would change every later block, so one symptom's
  classes would depend on what the case's other blocks hold; a case's blocks
  applied in another order (the later update wins) not reported; and the
  blocks HQ refuses for an empty case id keyed all as one, so a guard's block
  and an operation's, or the block HQ refused, are one class;
- the cases a refused submission leaves untouched reported in one class
  whatever the refusal's cause, so one register entry would hold two
  causes;
- the export paths left uncompared where HQ refuses Nova's next publish
  (decision 15 compares the local archive with ``build(A)`` always);
- a case attachment processed without the files a device sends beside its
  submission (HQ refuses each such run), or a run HQ refuses for want of
  them compared as equal on both sides;
- a soft assertion HQ notes while it processes a submission lost, rather
  than recorded with the operation that made it and held by proof 3;
- a case database the document's own filters and mappings find nothing in,
  or whose values sort the same as text and as numbers, so an ordering
  defect (defect 10) cannot show;
- a restore Core does not read whole, or one whose cases HQ's case processing
  does not hold, so an update or an index to a restore case is refused;
- HQ's builds read lookup tables from a restore that does not serve them;
- a search's own query keys (``case_type``, ``_xpath_query``) in one class
  with its prompts' keys, so a register entry for one absorbs the other;
- profile properties compared where no runtime running both builds reads
  them (a target's URLs), or left out where one does; the compared set, and
  which runtime reads each, must be the manifest's.
"""

from __future__ import annotations

import base64
import copy
import io
import json
import shutil
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path

import pytest
from lxml import etree

from proof.checks import casedata, cases, observations, proof2, proof3
from proof.checks.corpus import Document
from proof.checks.differences import Difference
from proof.core.client import CoreRunnerError
from proof.hq.check import hq_check
from proof.hq.configuration import Configuration
from proof.observe import sessions
from proof.observe.identity import data_namespace
from proof.observe.record import Blobs

XFORMS = "http://www.w3.org/2002/xforms"
XHTML = "http://www.w3.org/1999/xhtml"
CASE_XMLNS = "http://commcarehq.org/case/transaction/v2"
WORKTREE = Path(__file__).resolve().parents[2]
ENTRIES = WORKTREE / "lib" / "commcare" / "surface" / "entries"


@pytest.mark.parametrize("document", cases.document_params() + cases.control_params("proof3"))
def test_the_same_sessions_behave_the_same_on_each_build(document, hq, core_runner, editor_driver):
    found = proof3.document_behavior(
        document, observations.records_for(document, core_runner, editor_driver=editor_driver)
    )
    cases.hold("proof3", document, found, cases.load_register(), configurations=sorted(document.exports))


def _observed_without_hooks(document, core_runner, names):
    """The document's records under the configurations ``names`` alone, observed without the observation hooks.

    What the tests below read (A, B, B aligned, proof 3's sessions over them
    and the local archive) is observed by the unit itself, each from its
    state's mark, which every hook leaves the unit at; the hooks' own work
    (proof 4's saves in HQ's editors, the intent and manifest reads) lands
    in records these tests never read. Nothing is read from or kept in the
    evidence store, whose keys name the hooks' records too.
    """
    from proof.observe import unit

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(unit, "HOOKS", ())
        return unit.observe_document(document, core_runner=core_runner, configurations=set(names))


def _behavior(document, export, observed, core_runner):
    """Proof 3 on ``observed``'s builds: observed in a unit of the export's configuration as a document's tree
    observes it (``proof.observe.sessions``), then judged as its records are."""
    blobs = Blobs()
    a_build = observed.a.build if observed.a is not None else None
    differs = False
    if a_build is not None and observed.b_aligned is not None and observed.b_aligned.files is not None:
        raw = proof2.build_differences(document.id, a_build, observed.b_aligned, observed.alignment, rules=())
        differs = any(difference.kind != "refused" for difference in raw)
    with hq_check(export.configuration.hq(), validate=core_runner.validate_form) as (unit, _):
        restore_a = sessions.hq_restore(
            unit, casedata.case_database(document.document), export.create.lookups, "restore-a"
        )
        recorded = sessions.observe(
            unit,
            document=document,
            export=export,
            a_build=a_build,
            b_aligned=observed.b_aligned,
            b_differs=differs,
            restore_a=restore_a,
            core_runner=core_runner,
            blobs=blobs,
        )
    return proof3.behavioral_equivalence(
        document.id, observed, recorded, blobs, has_local=document.local_ccz is not None
    )


# The negative controls -------------------------------------------------------


def _element_paths(root):
    paths = {}

    def walk(element, prefix):
        path = f"{prefix}/{etree.QName(element).localname}"
        paths.setdefault(path, []).append(element)
        for child in element.iterchildren(etree.Element):
            walk(child, path)

    walk(root, "")
    return paths


def _alter_answer_path(form_bytes, data_path):
    """The form with the question Core reports at ``data_path`` stored under a renamed instance node.

    The node is found by the path its element names make; every element that
    names the node by its whole path as the node it controls, binds or sets
    (``ref`` or ``nodeset``: the control, its bind, a ``setvalue``) follows the
    rename, and a form whose body has other than one control for it gives None.
    An empty node keeps the old name beside it, so expressions reading the
    question (a case update's ``calculate``, a condition) still resolve, as
    Core requires, and now read nothing: only runs that open this form change.
    """
    form = etree.fromstring(form_bytes)
    data = form.find(f"{{{XHTML}}}head/{{{XFORMS}}}model/{{{XFORMS}}}instance")[0]
    nodes = _element_paths(data).get(data_path, [])
    controls = [e for e in form.find(f"{{{XHTML}}}body").iter(etree.Element) if e.get("ref") == data_path]
    if len(nodes) != 1 or nodes[0] is data or len(nodes[0]) or len(controls) != 1:
        return None
    node = nodes[0]
    name = f"{etree.QName(node).localname}_altered"
    altered = f"{data_path.rpartition('/')[0]}/{name}"
    original = node.tag
    node.tag = etree.QName(etree.QName(node).namespace, name).text
    node.addnext(etree.Element(original))
    for element in form.iter(etree.Element):
        for attribute in ("ref", "nodeset"):
            if element.get(attribute) == data_path:
                element.set(attribute, altered)
    return etree.tostring(form, encoding="utf-8", xml_declaration=True), altered


def _form_steps(run):
    return [step for step in run["trace"] if step.get("screen") == "form"]


def _case_update_reads(form_bytes):
    """The answer paths a case block's update reads whole: each bind on a case update's child whose calculate
    is exactly one answer's path, both found by the paths their element names make."""
    form = etree.fromstring(form_bytes)
    data = form.find(f"{{{XHTML}}}head/{{{XFORMS}}}model/{{{XFORMS}}}instance")[0]
    paths = _element_paths(data)
    reads = set()
    for bind in form.iter(f"{{{XFORMS}}}bind"):
        targets = paths.get(bind.get("nodeset"), [])
        update = targets[0].getparent() if len(targets) == 1 else None
        if update is not None and update.tag == f"{{{CASE_XMLNS}}}update" and bind.get("calculate") in paths:
            reads.add(bind.get("calculate"))
    return reads


def _xmlns(content):
    try:
        return data_namespace(content)
    except ValueError:
        return None


@dataclass
class Control:
    """A corpus document's build of A, its sessions, and the same build with one answer path altered."""

    document: Document
    export: object
    observed: observations.Republish
    restore: bytes
    baseline: proof3.Run
    # A's files with the altered form.
    files: dict
    original_path: str
    altered_path: str
    # The runs that answer the altered question (each must change) and that open its form (no other may).
    reaching: set
    opening: set


def _altered(observed, baseline):
    """A's build with one answer a case update reads stored under another path, in one submitted form."""
    files = observed.a.build.files
    for run in baseline.trace["runs"]:
        if run["end"] != "submitted":
            continue
        form = _form_steps(run)[-1]
        name = next(
            (key for key, content in files.items() if key.endswith(".xml") and _xmlns(content) == form["xmlns"]),
            None,
        )
        if name is None:
            continue
        reads = _case_update_reads(files[name])
        for event in form.get("events", []):
            if event.get("event") != "question" or event.get("dataPath") not in reads:
                continue
            result = _alter_answer_path(files[name], event["dataPath"])
            if result is None:
                continue
            opening, reaching = set(), set()
            for index, other in enumerate(baseline.trace["runs"]):
                for step in _form_steps(other):
                    if step["xmlns"] == form["xmlns"]:
                        opening.add(index)
                        if event["dataPath"] in [e.get("dataPath") for e in step.get("events", [])]:
                            reaching.add(index)
            return {**files, name: result[0]}, event["dataPath"], result[1], reaching, opening
    return None


@pytest.fixture(scope="module")
def control(hq, core_runner, tmp_path_factory):
    """The first corpus document whose build of A runs a session that submits an answer a case update reads.

    The document's publish sends no lookup tables, so its build of A, standing
    as a local archive, reads none from a restore that serves none (the local
    route's restore).
    """
    scratch = tmp_path_factory.mktemp("proof3-control")
    for document in cases.load_corpus().emitted:
        export = document.exports.get("minimum")
        if export is None or export.create.lookups is not None or export.republish.lookups is not None:
            continue
        records = _observed_without_hooks(document, core_runner, [export.configuration.name])
        observed = observations.republish_view(records, export.configuration.name)
        if (
            observed.a is None
            or observed.b_aligned is None
            or proof3.unbuildable(observed.a.build) is not None
            or document.local_ccz is None
        ):
            continue
        with hq_check(export.configuration.hq(), validate=core_runner.validate_form) as (state, _):
            restore = casedata.restore(casedata.case_database(document.document), state.domain)
        baseline = proof3.run_sessions(
            core_runner, "A", proof3.arrange_build(observed.a.build.files, scratch / document.id), restore
        )
        if baseline.trace is not None and (altered := _altered(observed, baseline)):
            return Control(document, export, observed, restore, baseline, *altered)
    pytest.fail(
        "No corpus document without lookup tables submits a form with an answer a case update reads that the"
        " control can alter, so neither route's negative control has an input."
    )


def _assert_reaching_runs_differ(found, control, artifacts):
    assert {d.artifact for d in found} == set(artifacts), [d.describe() for d in found]
    assert all(d.at.startswith("/runs/") for d in found), [d.describe() for d in found if not d.at.startswith("/runs/")]
    for artifact in artifacts:
        runs = {int(d.at.split("/")[2]) for d in found if d.artifact == artifact}
        assert control.reaching <= runs <= control.opening, (artifact, runs, control.reaching, control.opening)
    shown = [value for d in found for value in (d.before, d.after)]
    assert control.altered_path in shown and control.original_path in shown


def test_one_altered_answer_path_changes_the_trace_and_case_processing_of_the_runs_that_reach_it(
    control, core_runner, tmp_path
):
    """The plan's row "The Core runner's traces are faithful", on the B route, through proof 3's own comparison.

    A corpus document's build of A, and the same build with one answer a case
    update reads stored under another path in one submitted form, standing as
    B aligned: proof 2 sees the difference, so proof 3 runs B's sessions, and
    both its trace and HQ's case processing differ in every run that answers
    the question and in no run that does not open the form. The same build as
    B gives proof 2 nothing, so proof 3 runs no B session, and the same build's
    sessions twice compare equal.
    """
    document, export, observed = control.document, control.export, control.observed

    unchanged = copy.copy(observed)
    unchanged.b_aligned = replace(observed.a.build, state="B-aligned")
    assert proof2.build_equivalence(document.id, unchanged) == []
    assert [d for d in _behavior(document, export, unchanged, core_runner) if "@B" in d.artifact] == []

    again = proof3.run_sessions(
        core_runner,
        "B",
        proof3.arrange_build(observed.a.build.files, tmp_path / "again"),
        control.restore,
        proof3.script_of(control.baseline.trace),
    )
    assert (
        proof3.compare_traces(
            document.id, "trace@B", control.baseline.trace, again.trace, runtimes=proof3.RUNTIMES["B"]
        )
        == []
    )

    altered = copy.copy(observed)
    altered.b_aligned = replace(observed.a.build, state="B-aligned", files=control.files)
    assert proof2.build_equivalence(document.id, altered), "proof 2 must see the altered form"
    found = [d for d in _behavior(document, export, altered, core_runner) if "@B" in d.artifact]
    _assert_reaching_runs_differ(found, control, ("trace@B", "case_blocks@B"))


def _moved(element, old, new):
    """A copy of ``element`` and everything under it with the namespace ``old`` moved to ``new``."""
    name = etree.QName(element)
    moved = etree.Element(
        etree.QName(new, name.localname).text if name.namespace == old else element.tag,
        attrib=dict(element.attrib),
        nsmap={prefix: (new if uri == old else uri) for prefix, uri in element.nsmap.items()},
    )
    moved.text, moved.tail = element.text, element.tail
    for child in element:
        moved.append(_moved(child, old, new) if isinstance(child.tag, str) else copy.deepcopy(child))
    return moved


def _renamespaced(files):
    """HQ's build files with every form's ``xmlns`` replaced: in its data node and in the suite entry naming it."""
    renamed, mapping = dict(files), {}
    for name, content in sorted(files.items()):
        old = _xmlns(content) if name.endswith(".xml") else None
        if old is None:
            continue
        new = f"http://example.com/proof3/renamed/{len(mapping)}"
        mapping[old] = new
        form = etree.fromstring(content)
        data = form.find(f"{{{XHTML}}}head/{{{XFORMS}}}model/{{{XFORMS}}}instance")[0]
        data.getparent().replace(data, _moved(data, old, new))
        renamed[name] = etree.tostring(form, encoding="utf-8", xml_declaration=True)
    suite = etree.fromstring(files["suite.xml"])
    for form in suite.iter("form"):
        if form.getparent().tag == "entry" and (form.text or "").strip() in mapping:
            form.text = mapping[form.text.strip()]
    renamed["suite.xml"] = etree.tostring(suite, encoding="utf-8", xml_declaration=True)
    return renamed, mapping


def _as_local_archive(control, files, directory):
    """The control's document carrying ``files`` (HQ's build files) as its local archive, as HQ arranges them."""
    arranged = proof3.arrange_build(files, directory / "arranged")
    root = directory / "document"
    root.mkdir(parents=True)
    shutil.copyfile(control.document.document_path, root / "document.json")
    with zipfile.ZipFile(root / "local.ccz", "w") as archive:
        for path in sorted(p for p in arranged.rglob("*") if p.is_file()):
            archive.write(path, path.relative_to(arranged).as_posix())
    return Document(id=control.document.id, source=control.document.source, root=root)


def test_one_altered_answer_path_changes_the_local_archives_runs_that_reach_it(control, core_runner, tmp_path):
    """The plan's row "The Core runner's traces are faithful", on the local route proof 3 always runs.

    A's build with every form's ``xmlns`` replaced stands as the document's
    local archive. Its sessions compare equal with A's, so the namespace
    mapping (proof 1 judges identity) maps everything the trace and the
    submission carry, and the local restore serves what A's does. The same
    archive with the answer path altered differs, in its trace and in HQ's
    case processing, in every run that answers the question and in no run
    that does not open the form.
    """
    local_only = copy.copy(control.observed)
    local_only.b_aligned = None

    renamed, mapping = _renamespaced(control.observed.a.build.files)
    unchanged = _as_local_archive(control, renamed, tmp_path / "unchanged")
    with sessions.admitted(core_runner, unchanged.local_ccz) as report:
        alignment = proof3.xmlns_alignment(report, control.baseline.admission)
    assert alignment and alignment == {new: old for old, new in mapping.items()}
    assert _behavior(unchanged, control.export, local_only, core_runner) == []

    altered = _as_local_archive(control, _renamespaced(control.files)[0], tmp_path / "altered")
    found = _behavior(altered, control.export, local_only, core_runner)
    _assert_reaching_runs_differ(found, control, ("trace@local.ccz", "case_blocks@local.ccz"))


def test_no_session_runs_on_a_build_hq_refuses_to_make(control, core_runner):
    """A build whose ``validate_app`` lists an error is no archive a runtime installs: no session runs on it.

    The control's own builds run sessions (the tests above); with an error
    HQ's validation lists, A's stops the whole document before any session,
    which the bar reports, and B's refuses B's comparison alone, its cause
    the error's path.
    """
    error = {"type": "form filter has xpath error", "message": "Expecting 'QNAME', got 'AT'"}
    refused_a = copy.copy(control.observed)
    refused_a.a = replace(control.observed.a, build=replace(control.observed.a.build, errors=[error]))
    # The bar reports why A builds nothing a runtime installs (``validate_app@A``), so proof 3 runs no session
    # and reports nothing more.
    assert _behavior(control.document, control.export, refused_a, core_runner) == []

    refused_b = copy.copy(control.observed)
    refused_b.b_aligned = replace(control.observed.a.build, state="B-aligned", files=control.files, errors=[error])
    found = _behavior(control.document, control.export, refused_b, core_runner)
    cause = "/unbuildable/validate_app/form filter has xpath error"
    assert [d for d in found if "@B" in d.artifact] == [
        Difference("proof3", control.document.id, "trace@B", cause, cause, "refused", None, error)
    ]


# The case database -----------------------------------------------------------


def _document():
    """A document as ``document.json`` holds one: a parent and an extension type, a mapped column, a filter."""
    prose = {"parts": [{"kind": "text", "text": "Label"}]}
    term = lambda value: {"kind": "term", "term": value}  # noqa: E731
    return {
        "doc": {
            "caseTypes": [
                {
                    "name": "visit",
                    "parent_type": "patient",
                    "relationship": "extension",
                    "properties": [{"name": "score", "label": prose, "data_type": "int"}],
                },
                {
                    "name": "patient",
                    "properties": [
                        {"name": "case_name", "label": prose},
                        {"name": "code", "label": prose},
                        {"name": "weight", "label": prose, "data_type": "decimal"},
                        {
                            "name": "colour",
                            "label": prose,
                            "data_type": "single_select",
                            "options": [{"value": "red", "label": prose}, {"value": "blue", "label": prose}],
                        },
                    ],
                },
            ],
            "moduleOrder": ["m"],
            "modules": {
                "m": {
                    "caseType": "patient",
                    "caseListConfig": {
                        "columns": [
                            {"kind": "id-mapping", "field": "code", "mapping": [{"value": "a"}, {"value": "b"}]}
                        ],
                        "filter": {
                            "kind": "gt",
                            "left": term({"kind": "prop", "caseType": "patient", "property": "weight"}),
                            "right": term({"kind": "literal", "value": 40}),
                        },
                    },
                    "caseSearchConfig": {
                        "excludedOwnerIds": term({"kind": "session-user", "field": "excluded_owners"}),
                    },
                }
            },
            "userProperties": {"u": {"uuid": "u", "slug": "district"}},
        }
    }


def test_the_case_database_holds_what_the_documents_lists_select_in_orders_that_differ():
    database = casedata.case_database(_document())
    assert database == casedata.case_database(_document())
    by_type = {}
    for case in database.cases:
        by_type.setdefault(case.case_type, []).append(case)
    assert list(by_type) == ["patient", "visit", casedata.USERCASE_TYPE]
    patients = by_type["patient"]
    values = {name: [dict(case.properties).get(name) for case in patients] for name in ("code", "weight", "colour")}
    # The document's own mapping values and filter literal are held, beside values ordered three ways.
    assert {"a", "b", "10", "2"} <= set(values["code"])
    assert {"40", "10.5", "2.25"} <= set(values["weight"])
    assert {"red", "blue"} <= set(values["colour"])
    names = [case.name for case in patients]
    stored = names[:3]
    assert (
        stored != sorted(stored)
        and sorted(stored, key=lambda v: (not v.isdigit(), int(v) if v.isdigit() else 0)) != stored
    )
    scores = [dict(case.properties)["score"] for case in by_type["visit"]]
    assert scores != sorted(scores) and scores != sorted(scores, key=int)
    assert all(
        case.indices == (("parent", "patient", patients[n % len(patients)].case_id, "extension"),)
        for n, case in enumerate(by_type["visit"])
    )
    assert dict(database.user_data) == {"district": "proof district", "excluded_owners": "proof excluded_owners"}
    assert all(case.owner_id == casedata.USER_ID for case in database.cases)


def _casedb_form():
    return (
        '<h:html xmlns="http://www.w3.org/2002/xforms" xmlns:h="http://www.w3.org/1999/xhtml">'
        "<h:head><h:title>Cases</h:title><model>"
        '<instance><data xmlns="http://example.org/cases"><n/></data></instance>'
        '<instance id="casedb" src="jr://instance/casedb"/>'
        '<instance id="commcaresession" src="jr://instance/session"/>'
        '<bind nodeset="/data/n"/>'
        '</model></h:head><h:body><input ref="/data/n"><label>n</label></input></h:body></h:html>'
    ).encode()


def test_core_reads_the_whole_restore_hq_writes(hq, core_runner):
    database = casedata.case_database(_document())
    with hq_check(Configuration(), validate=core_runner.validate_form) as (state, _):
        restore = casedata.restore(database, state.domain)
        assert restore == casedata.restore(database, state.domain)
    result = core_runner.evaluate(
        formBase64=base64.b64encode(_casedb_form()).decode(),
        restoreBase64=base64.b64encode(restore).decode(),
        expressions=[
            "instance('casedb')/casedb/case[@case_type='visit'][1]/index/parent",
            "instance('commcaresession')/session/user/data/district",
        ],
        instances=["casedb"],
    )
    casedb = etree.fromstring(result["instance"]["casedb"].encode())
    # Each case as Core's casedb holds it: its type, its name and the restore's values of its properties
    # (Core adds its own elements beside them, such as attachment and index).
    held = {
        case.get("case_id"): (case.get("case_type"), case.findtext("case_name"), case) for case in casedb.iter("case")
    }
    assert set(held) == {case.case_id for case in database.cases}
    for case in database.cases:
        case_type, name, element = held[case.case_id]
        expected = {**dict(case.properties), **({"external_id": case.external_id} if case.external_id else {})}
        assert (case_type, name) == (case.case_type, case.name)
        assert {key: element.findtext(key) for key in expected} == expected
    first_visit = next(case for case in database.cases if case.case_type == "visit")
    assert [value["value"] for value in result["values"]] == [first_visit.indices[0][2], "proof district"]


# HQ's case processing --------------------------------------------------------


def _submission(*blocks):
    body = "".join(blocks)
    return f"""<?xml version='1.0' ?>
<data xmlns="http://example.com/proof/proof3" name="Visit">{body}
  <meta xmlns="http://openrosa.org/jr/xforms">
    <instanceID>form-proof3</instanceID><userID>{casedata.USER_ID}</userID>
    <timeStart>2026-01-15T10:30:00.000Z</timeStart><timeEnd>2026-01-15T10:30:00.000Z</timeEnd>
  </meta>
</data>""".encode()


def _block(case_id, inner):
    return (
        f'<case xmlns="{CASE_XMLNS}" case_id="{case_id}" date_modified="2026-01-15T10:30:00.000Z"'
        f' user_id="{casedata.USER_ID}">{inner}</case>'
    )


def test_hq_processes_a_submission_over_the_restores_cases(hq, core_runner):
    """An update to a restore case applies over its restored properties, and an index to one resolves; without
    the restore's cases HQ holds only what the submission says, and refuses the index."""
    from casexml.apps.case.exceptions import InvalidCaseIndex

    from proof.hq import operations

    database = casedata.case_database(_document())
    patient = next(case for case in database.cases if case.case_type == "patient")
    update = _submission(_block(patient.case_id, "<update><code>b</code></update>"))
    child = _submission(
        _block(
            "visit-new",
            "<create><case_type>visit</case_type><case_name>New</case_name>"
            f"<owner_id>{casedata.USER_ID}</owner_id></create>"
            f'<index><parent case_type="patient" relationship="extension">{patient.case_id}</parent></index>',
        )
    )
    with hq_check(Configuration(), validate=core_runner.validate_form) as (state, _):
        updated = operations.process_case_blocks(state, update, cases=casedata.hq_cases(database, state.domain))
        indexed = operations.process_case_blocks(state, child, cases=casedata.hq_cases(database, state.domain))
        bare_update = operations.process_case_blocks(state, update)
        bare_index = operations.process_case_blocks(state, child)
    assert updated.refusal is None and indexed.refusal is None
    assert updated.touched[patient.case_id].case.case_json == {**dict(patient.properties), "code": "b"}
    assert not updated.touched[patient.case_id].is_creation
    assert [(i.referenced_id, i.relationship) for i in indexed.touched["visit-new"].case.live_indices] == [
        (patient.case_id, "extension")
    ]
    assert bare_update.touched[patient.case_id].case.case_json == {"code": "b"}
    assert isinstance(bare_index.refusal, InvalidCaseIndex)


def test_hq_names_where_the_form_holds_each_case_block_it_reads(hq, core_runner):
    """The observation records HQ's own path to each block it reads (``extract_case_blocks`` with
    ``include_path``), so where HQ refuses a submission for an empty case id (``IllegalCaseId``), a guard's block
    and an operation's are told apart by where the form holds them, and so is the block HQ refused. Reading the
    paths notes nothing HQ's processing did not: a guard that lost its ``date_modified`` too is noted as often as
    HQ's processing alone notes it."""
    from proof.hq import operations
    from proof.observe.unit import OperationLog

    database = casedata.case_database(_document())
    patient = next(case for case in database.cases if case.case_type == "patient")
    update = _block(patient.case_id, "<update><code>b</code></update>")
    guard = _block("", "<update><visit_status>a</visit_status></update>").replace(
        'date_modified="2026-01-15T10:30:00.000Z"', 'date_modified=""'
    )
    promote = _block(
        "",
        f"<create><case_type>visit</case_type><case_name>n</case_name><owner_id>{casedata.USER_ID}</owner_id></create>",
    )
    held = f"<{GUARD[0]}><{GUARD[1]}>{guard}</{GUARD[1]}><{OPERATION[1]}>{promote}</{OPERATION[1]}></{GUARD[0]}>"

    def run(*blocks):
        return {"trace": [{"screen": "form", "submission": _submission(*blocks).decode()}], "generated": {}}

    with hq_check(Configuration(), validate=core_runner.validate_form) as (state, record):
        log = OperationLog(state, record, {})
        accepted = sessions.processed_runs(state, database, {"runs": [run(update)]})
        with log("case-processing:A", b"hq"):
            operations.process_case_blocks(
                state, _submission(update, held), cases=casedata.hq_cases(database, state.domain)
            )
        with log("case-processing:B", b"observed"):
            refused = sessions.processed_runs(state, database, {"runs": [run(update, held)]})
    (processed,) = refused["runs"]
    assert processed["refusal"] == {"class": "IllegalCaseId", "message": "case_id must not be empty"}
    assert processed["blockPaths"] == [[], GUARD, OPERATION]
    assert [block["@case_id"] for block in processed["blocks"]] == [patient.case_id, "", ""]
    assert accepted["runs"][0]["blockPaths"] == [[]]
    hq_alone, observed = log.log
    assert hq_alone["softAssertions"]
    assert [(note["where"], note["message"]) for note in observed["softAssertions"]] == [
        (note["where"], note["message"]) for note in hq_alone["softAssertions"]
    ]
    found = proof3.compare_case_processing("d", "case_blocks@B", accepted, refused)
    assert {(d.path, d.kind) for d in found} == {
        (GUARD_KEY, "added"),
        (OPERATION_KEY, "added"),
        ("/runs/*/refusal/IllegalCaseId/create/__nova_operations~1*~1case", "added"),
        ("/runs/*/refusal/IllegalCaseId/create/__nova_operations~1*~1case/cases/*", "removed"),
    }


def test_a_case_attachment_is_read_from_the_files_the_device_sends_and_without_them_is_refused(hq, core_runner):
    """HQ reads a case attachment's file from the form's files only under ``MM_CASE_PROPERTIES``
    (``SqlCaseUpdateStrategy._apply_attachments_action``), from the files stored on the form as it is read
    (``_create_new_xform``). Given a stand-in for each file its case blocks name, as a device sends them, the run
    is processed and compares equal; without them it is a refused comparison on both sides, never an equal one.
    Without the flag HQ ignores the attachment either way."""
    database = casedata.case_database(_document())
    patient = next(case for case in database.cases if case.case_type == "patient")
    run = {
        "trace": [
            {
                "screen": "form",
                "submission": _submission(
                    _block(patient.case_id, '<attachment><photo src="photo.jpg" from="local"/></attachment>')
                ).decode(),
            }
        ],
        "generated": {"clock": {}, "uuids": 0},
    }
    processed = {}
    for name, configuration in (
        ("with", Configuration(flags={"MM_CASE_PROPERTIES"})),
        ("without", Configuration()),
    ):
        with hq_check(configuration, validate=core_runner.validate_form) as (state, _):
            processed[name] = sessions.processed_runs(state, database, {"runs": [run]})
            processed[f"{name}, no files"] = sessions.processed_runs(state, database, {"runs": [run]}, files=False)
    assert sessions.stand_ins(run["trace"][0]["submission"]) == {"photo.jpg": (sessions.STAND_IN, "image/jpeg")}
    sent = processed["with"]["runs"][0]
    assert "needsFiles" not in sent and sent["cases"][patient.case_id]["creation"] is False
    assert proof3.compare_case_processing("d", "case_blocks@B", processed["with"], processed["with"]) == []

    needing = processed["with, no files"]["runs"][0]["needsFiles"]
    assert needing["class"] == "AttachmentNotFound" and "photo.jpg" in needing["message"]
    assert proof3.compare_case_processing(
        "d", "case_blocks@B", processed["with, no files"], processed["with, no files"]
    ) == [
        Difference(
            "proof3",
            "d",
            "case_blocks@B",
            "/runs/*/needs-files/AttachmentNotFound",
            "/runs/0/needs-files/AttachmentNotFound",
            "refused",
            needing,
            needing,
        )
    ]

    for name in ("without", "without, no files"):
        assert "needsFiles" not in processed[name]["runs"][0]
        assert processed[name]["runs"][0]["cases"][patient.case_id]["creation"] is False
    assert processed["without"] == processed["without, no files"]


def test_a_corpus_document_whose_forms_send_files_is_processed_with_them(hq, core_runner):
    """A case-extension document's sessions submit a photo its case block attaches: under ``MM_CASE_PROPERTIES``,
    with the files the device sends, HQ processes every run, where without them HQ refuses each one
    (``needsFiles``)."""
    document = next((d for d in cases.load_corpus().emitted if d.id.startswith("case-extension") and d.local_ccz), None)
    assert document is not None, "The corpus holds no case-extension document, whose forms attach a file."
    name = next(
        (
            name
            for name, export in sorted(document.exports.items())
            if "MM_CASE_PROPERTIES" in export.configuration.flags
        ),
        None,
    )
    assert name is not None, f"{document.id} is exported under no configuration where HQ reads case attachments."
    records = _observed_without_hooks(document, core_runner, [name])
    recorded = observations.sessions_record(records, name)
    trace = records.blobs.get_json(recorded["local"]["trace"])
    sent = records.blobs.get_json(recorded["local"]["processed"])
    attaching = [
        index
        for index, run in enumerate(trace["runs"])
        if proof3.submission_of(run) and sessions.stand_ins(proof3.submission_of(run))
    ]
    assert attaching, f"{document.id}'s sessions submit no form whose case blocks attach a file."
    assert all("needsFiles" not in sent["runs"][index] for index in attaching)
    database = casedata.case_database(document.document)
    xmlns = proof3.xmlns_alignment(recorded["local"]["admission"], recorded["baseline"]["admission"])
    with hq_check(document.exports[name].configuration.hq(), validate=core_runner.validate_form) as (state, _):
        bare = sessions.processed_runs(state, database, trace, xmlns, files=False)
    assert all(bare["runs"][index]["needsFiles"]["class"] == "AttachmentNotFound" for index in attaching)


def test_each_note_hq_makes_processing_a_submission_is_recorded_with_its_operation_and_held_by_proof_3(hq, core_runner):
    """An empty ``date_modified`` is a soft assertion HQ notes and goes on from, as production does
    (``proof.hq.boot``): the operation that processed the submission records the note in its part, and proof 3
    reports it as its own difference; the same submission dated notes nothing."""
    from proof.hq import operations
    from proof.observe.record import ConfigurationRecords, DocumentRecords
    from proof.observe.unit import OperationLog

    undated = _submission(
        _block("visit-1", "<update><code>b</code></update>").replace(
            'date_modified="2026-01-15T10:30:00.000Z"', 'date_modified=""'
        )
    )
    dated = _submission(_block("visit-1", "<update><code>b</code></update>"))
    with hq_check(Configuration(), validate=core_runner.validate_form) as (state, record):
        log = OperationLog(state, record, {})
        with log("case-processing:local.ccz", b"undated"):
            operations.process_case_blocks(state, undated)
        with log("case-processing:A", b"dated"):
            operations.process_case_blocks(state, dated)
    noted, quiet = log.log
    assert noted["softAssertions"] and all(
        note["where"].endswith("casexml/apps/case/util.py::validate_phone_datetime") for note in noted["softAssertions"]
    )
    assert "softAssertions" not in quiet
    records = DocumentRecords("d", "corpus")
    records.configurations["minimum"] = ConfigurationRecords("minimum", {}, b_aligned={"operations": log.log})
    found = observations.soft_assertion_differences(records, "proof3")
    assert {(d.artifact, d.kind) for d in found} == {("soft_assert:case-processing:local.ccz@B", "error")}
    assert observations.soft_assertion_differences(records, "bar") == []


def test_each_note_a_hooks_operation_or_request_makes_is_recorded_under_the_hook(hq, core_runner):
    """What an observation hook runs through its unit (``proof.observe.unit.HookUnit``) is logged in its part's
    record under the hook's name: each operation, and each request in which HQ noted something; and the hook's
    check holds the notes, no other."""
    from proof.hq import operations
    from proof.observe.record import ConfigurationRecords, DocumentRecords
    from proof.observe.unit import HookUnit, OperationLog

    undated = _submission(
        _block("visit-1", "<update><code>b</code></update>").replace(
            'date_modified="2026-01-15T10:30:00.000Z"', 'date_modified=""'
        )
    )
    dated = _submission(_block("visit-1", "<update><code>b</code></update>"))
    with hq_check(Configuration(), validate=core_runner.validate_form) as (state, record):
        log = OperationLog(state, record, {})
        hook = HookUnit(state, log, "proof4")
        with hook.operation("probe", b"undated"):
            operations.process_case_blocks(state, undated)
        with hook.request(b"\x01" * 32):
            operations.process_case_blocks(state, undated)
        with hook.request(b"\x02" * 32):
            operations.process_case_blocks(state, dated)
    (operation,) = log.log
    assert operation["hook"] == "proof4" and operation["softAssertions"]
    (request,) = log.requests
    assert request["hook"] == "proof4" and request["request"] == "01" * 32 and request["softAssertions"]
    records = DocumentRecords("d", "corpus")
    records.configurations["minimum"] = ConfigurationRecords("minimum", {}, b=log.recorded({}))
    found = {(d.artifact, d.kind) for d in observations.soft_assertion_differences(records, "proof4")}
    assert found == {("soft_assert:probe@B", "error"), (f"soft_assert:request:{'01' * 8}@B", "error")}
    assert observations.soft_assertion_differences(records, "bar") == []


# A republish HQ refuses -------------------------------------------------------


def _with_app_file(captured, app_json):
    """The captured upload's bytes with its ``app_file`` field holding ``app_json``, every other byte as captured."""
    from proof.hq.operations import _boundary, _form_parts, _part_name

    boundary = _boundary(captured.content_type)
    preamble, parts, epilogue = _form_parts(captured.body_path.read_bytes(), boundary)
    rebuilt = [preamble]
    for head, content in parts:
        if _part_name(head) == "app_file":
            content = json.dumps(app_json).encode("utf-8")
        rebuilt.append(b"\r\n" + head + b"\r\n\r\n" + content + b"\r\n")
    rebuilt.append(epilogue)
    return (b"--" + boundary).join(rebuilt)


def _submitting_subject(core_runner):
    """The cheapest document with a local archive whose sessions on it submit a form, its configuration and records
    (under that configuration)."""
    from proof.checks import sharding

    estimate = sharding.estimate(sharding.load_timings())
    documents = sorted(
        (document for document in cases.load_corpus().emitted if document.local_ccz is not None),
        key=lambda document: (estimate(document.group), document.id),
    )
    for document in documents:
        for name in sorted(document.exports):
            records = _observed_without_hooks(document, core_runner, [name])
            recorded = observations.sessions_record(records, name) or {}
            processed = (recorded.get("local") or {}).get("processed")
            if processed and any(run["submitted"] for run in records.blobs.get_json(processed)["runs"]):
                return document, name, records
    raise AssertionError("No corpus document's sessions submit a form on its local archive, so nothing is shown.")


def _local_route(differences):
    return sorted(json.dumps(d.as_json(), sort_keys=True) for d in differences if d.artifact.endswith("@local.ccz"))


def test_a_republish_hq_refuses_still_compares_the_local_archive_with_build_a(hq, core_runner, tmp_path):
    """The local archive is compared with ``build(A)`` always (decision 15), whatever HQ makes of the next publish:
    where HQ refuses Nova's republish (its app file no app), the ``b_aligned`` record still holds the sessions of A
    and of the local archive, from A's state, and proof 3 judges them as it does where HQ accepts it; there is no
    B to align, so no B route."""
    document, name, held = _submitting_subject(core_runner)
    accepted = proof3.behavioral_equivalence(
        document.id,
        observations.republish_view(held, name),
        observations.sessions_record(held, name),
        held.blobs,
        has_local=True,
    )
    root = tmp_path / document.id
    shutil.copytree(document.root, root)
    (root / "inputs.json").unlink()  # the copy's keys are its own files' (unit.computed_inputs)
    refusing = Document(id=document.id, source=document.source, root=root)
    republish = refusing.exports[name].republish
    republish.body_path.write_bytes(_with_app_file(republish, {}))

    records = _observed_without_hooks(refusing, core_runner, [name])
    observed = records.configurations[name]
    assert observed.b["publish"]["refused"] is not None and "state" not in observed.b
    assert "aligned" not in observed.b_aligned
    recorded = observed.b_aligned.get("sessions")
    assert recorded is not None, "HQ refused B, and the b_aligned record holds no sessions of A or the local archive."
    assert recorded["local"]["trace"] and recorded["local"]["processed"] and "b" not in recorded
    refused = proof3.behavioral_equivalence(
        refusing.id,
        observations.republish_view(records, name),
        observations.sessions_record(records, name),
        records.blobs,
        has_local=True,
    )
    assert _local_route(refused) == _local_route(accepted)
    assert not [d for d in refused if d.artifact.endswith("@B")]


# Lookup tables ---------------------------------------------------------------


def _lookup_document():
    for document in cases.load_corpus().emitted:
        for export in document.exports.values():
            if export.create.lookups is not None:
                return document, export
    return None


def test_hqs_builds_read_lookup_tables_from_hqs_restore(hq, core_runner):
    found = _lookup_document()
    assert found, "No corpus document's publish sends lookup tables, so nothing shows HQ's restore serves them."
    document, export = found
    lookup = document.document["lookup"]
    definition = lookup["definitions"][0]
    rows = lookup["rowsByTable"][definition["id"]]
    column = definition["columns"][0]
    tag = definition["tag"]
    form = _casedb_form().replace(
        b'<instance id="casedb"',
        f'<instance id="item-list:{tag}" src="jr://fixture/item-list:{tag}"/><instance id="casedb"'.encode(),
    )
    database = casedata.case_database(document.document)
    with hq_check(export.configuration.hq(), validate=core_runner.validate_form) as (state, _):
        tables = casedata.lookup_fixtures(state, export.create.lookups)
        restore = casedata.restore(database, state.domain, tables)
        plain = casedata.restore(database, state.domain)
    expressions = [
        f"instance('item-list:{tag}')/{tag}_list/{tag}[{index + 1}]/{column['wireName']}" for index in range(len(rows))
    ]
    result = core_runner.evaluate(
        formBase64=base64.b64encode(form).decode(),
        restoreBase64=base64.b64encode(restore).decode(),
        expressions=[f"count(instance('item-list:{tag}')/{tag}_list/{tag})", *expressions],
    )
    values = [value["value"] for value in result["values"]]
    assert values == [str(len(rows)), *(row["values"][column["id"]] for row in rows)]
    with pytest.raises(CoreRunnerError, match=f"Unable to find lookup table: jr://fixture/item-list:{tag}"):
        core_runner.evaluate(
            formBase64=base64.b64encode(form).decode(),
            restoreBase64=base64.b64encode(plain).decode(),
            expressions=[f"count(instance('item-list:{tag}')/{tag}_list/{tag})"],
        )


@dataclass(frozen=True)
class _Workbook:
    """A lookup workbook capture as ``casedata.lookup_fixtures`` reads one."""

    body_path: Path
    content: bytes

    def workbook(self):
        return self.content


def _per_owner(workbook):
    """The workbook with every table on its types sheet marked as held per owner (``is_global?`` ``no``)."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(workbook))
    sheet = book["types"]
    header = [cell.value for cell in sheet[1]]
    column = header.index("is_global?") + 1
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=column, value="no")
    written = io.BytesIO()
    book.save(written)
    return written.getvalue()


def test_a_lookup_table_hq_cannot_serve_is_a_refused_comparison(hq, core_runner):
    """A table HQ holds per owner is served by no restore here: proof 3 reports one refused comparison naming
    it, which the register can hold, instead of raising. The same document's global tables are served (the test
    above)."""
    found = _lookup_document()
    assert found, "No corpus document's publish sends lookup tables, so nothing shows a refused table."
    document, export = found
    name = export.configuration.name
    observed = observations.republish_view(_observed_without_hooks(document, core_runner, [name]), name)
    assert observed.a is not None and proof3.unbuildable(observed.a.build) is None
    captured = export.create.lookups
    per_owner = replace(
        export, create=replace(export.create, lookups=_Workbook(captured.body_path, _per_owner(captured.workbook())))
    )
    differences = _behavior(document, per_owner, observed, core_runner)
    assert [(d.artifact, d.path, d.kind) for d in differences] == [
        ("trace@A", "/lookups-not-served/per-owner-table", "refused")
    ]
    assert differences[0].after["workbook"] == captured.body_path.name
    assert differences[0].after["table"] in {table["tag"] for table in document.document["lookup"]["definitions"]}


# The profile and the search keys ----------------------------------------------


def _trace(properties, required=("2", "57", "0"), runs=()):
    return {
        "derived": True,
        "profile": {
            "requiredVersion": dict(zip(("requiredMajor", "requiredMinor", "requiredMinimal"), required, strict=True)),
            "properties": [{"key": key, "value": value, "force": True} for key, value in properties],
        },
        "runs": list(runs),
    }


def test_a_profile_is_compared_on_what_a_runtime_running_both_builds_reads(without_spelling_rules):
    hq_built = _trace([("ota-restore-url", "https://a"), ("cc-show-saved", "no"), ("cc-persistent-menu", "no")])
    local = _trace([("ota-restore-url", "https://b"), ("cc-persistent-menu", "yes")], required=(None, None, None))
    across_paths = proof3.compare_traces("d", "trace@local.ccz", hq_built, local, runtimes=proof3.RUNTIMES["local.ccz"])
    assert {d.path for d in across_paths} == {
        "/profile/properties/cc-show-saved",
        "/profile/requiredVersion/requiredMajor",
        "/profile/requiredVersion/requiredMinor",
        "/profile/requiredVersion/requiredMinimal",
    }
    between_builds = proof3.compare_traces("d", "trace@B", hq_built, local, runtimes=proof3.RUNTIMES["B"])
    assert "/profile/properties/cc-persistent-menu/value" in {d.path for d in between_builds}
    assert not [d for d in between_builds if "ota-restore-url" in d.path]


def _search(params, prompts=("name",)):
    """A trace of one run whose one step is a search with ``prompts``, sending ``params`` on its step and request."""
    step = {
        "screen": "search",
        "prompts": [{"key": key, "text": key} for key in prompts],
        "chosen": {"search": {key: "proof" for key in prompts}},
        "params": params,
        "promptErrors": {key: None for key in prompts},
        "requests": [{"kind": "search", "params": params}],
    }
    return _trace([], runs=[{"trace": [step], "script": [], "end": "search-failed"}])


def test_a_searchs_own_query_keys_and_its_prompts_keys_are_separate_classes():
    full = {"case_type": ["patient"], "_xpath_query": ["match-all()"], "name": ["proof"]}
    paths = {}
    for key in full:
        fewer = {name: value for name, value in full.items() if name != key}
        found = proof3.compare_traces("d", "trace@B", _search(full), _search(fewer), runtimes=proof3.RUNTIMES["B"])
        paths[key] = {d.path for d in found}
        assert {d.at.rpartition("/")[2] for d in found} == {key}
    assert paths == {
        "case_type": {"/runs/*/trace/*/params/case_type", "/runs/*/trace/*/requests/*/params/case_type"},
        "_xpath_query": {"/runs/*/trace/*/params/_xpath_query", "/runs/*/trace/*/requests/*/params/_xpath_query"},
        "name": {"/runs/*/trace/*/params/*", "/runs/*/trace/*/requests/*/params/*"},
    }
    # A prompt only one side asks is the app's data on that side, in its answers, errors and parameters alike.
    found = proof3.compare_traces(
        "d", "trace@B", _search(full), _search({**full, "age": ["7"]}, ("name", "age")), runtimes=proof3.RUNTIMES["B"]
    )
    assert {d.path for d in found if d.at.endswith("/age")} == {
        "/runs/*/trace/*/chosen/search/*",
        "/runs/*/trace/*/params/*",
        "/runs/*/trace/*/promptErrors/*",
        "/runs/*/trace/*/requests/*/params/*",
    }


def _profile_key(surface_key):
    """The profile property key HQ writes a setting under, or None when HQ writes it elsewhere."""
    from corehq.apps.app_manager.const import ANDROID_LOGO_PROPERTY_MAPPING

    family, _, name = surface_key.partition(":")
    kind, _, setting = name.partition(".")
    if family != "setting":
        return None
    if kind == "properties":
        return ANDROID_LOGO_PROPERTY_MAPPING.get(f"hq_{setting}", setting)
    # app_manager/templates/app_manager/profile.xml writes these two app settings as properties.
    return {"persistent_menu": "cc-persistent-menu", "show_breadcrumbs": "cc-breadcrumbs-enabled"}.get(setting)


RUNTIME_READINGS = ("runs", "different")


def test_the_compared_profile_properties_are_the_manifests_runtime_read_app_content(hq):
    """Each compared property is a HELD or HELD-NEW setting a runtime reads, read by exactly the runtimes the
    manifest says read it, where the built profile is what that runtime reads."""
    held = {}
    for path in sorted(ENTRIES.glob("*.json")):
        for entry in json.loads(path.read_text(encoding="utf-8")):
            platforms = entry.get("platforms") if isinstance(entry.get("platforms"), dict) else {}
            reads = {
                runtime: (platforms.get(runtime) or {}).get("reading") in RUNTIME_READINGS
                for runtime in ("android", "webApps")
            }
            if entry.get("disposition") not in ("HELD", "HELD-NEW") or not any(reads.values()):
                continue
            for key in entry.get("surfaceKeys", []):
                profile_key = _profile_key(key)
                if profile_key is not None:
                    held[profile_key] = (key, reads)
            if entry.get("valueClass") == "custom-property-index-case-search-results":
                held["cc-index-case-search-results"] = (entry["id"], reads)
    # Web Apps shows the logo HQ holds for the app (cloudcare/views.py::_format_app_doc, logo_refs); no runtime
    # reads the profile property HQ writes for it.
    assert held.pop("brand-banner-web-apps") == (
        "setting:properties.logo_web_apps",
        {"android": False, "webApps": True},
    )
    # Web Apps hides incomplete forms by the property of HQ's app document, not of the profile it installs
    # (cloudcare/static/cloudcare/js/formplayer/apps/controller.js, isIncompleteFormsDisabled, and apps/views.js
    # through getAppDisplayProperties, which app.js sets from the app document cloudcare/views.py::_format_app_doc
    # serves), so the built profile's property has an Android reader alone.
    setting, reads = held["cc-show-incomplete"]
    assert reads == {"android": True, "webApps": True}
    held["cc-show-incomplete"] = (setting, {"android": True, "webApps": False})
    assert {key: setting for key, (setting, _) in held.items()} == {
        key: readers.setting for key, readers in proof3.PROFILE_PROPERTIES.items()
    }
    for key, (_, reads) in held.items():
        readers = proof3.PROFILE_PROPERTIES[key]
        assert {"android": readers.android is not None, "webApps": readers.web_apps is not None} == reads, key


# HQ's case processing, keyed as HQ applies it, and refusals by cause ----------------------------------------


def _applied_block(case_id, **actions):
    return {"@case_id": case_id, "@date_modified": "now", "@user_id": "w", **actions}


def _processed_record(*blocks, refusal=None, cases=None, paths=None):
    """One run of HQ's case processing as the observation records it, each block at the form's root unless
    ``paths`` names where the form holds each (HQ's path to it, ``sessions.block_paths``)."""
    return {
        "runs": [
            {
                "submitted": True,
                "blocks": list(blocks),
                "blockPaths": [list(path) for path in paths] if paths is not None else [[] for _ in blocks],
                "refusal": refusal,
                "cases": cases or {},
            }
        ]
    }


def _case_classes(before, after, traces=None):
    found = proof3.compare_case_processing("d", "case_blocks@B", before, after, traces=traces)
    return {(d.path, d.kind) for d in found}


def test_case_blocks_are_compared_as_hq_applies_them():
    """HQ applies a submission's blocks by case id, each case's by its first action in ``CASE_ACTIONS`` and then
    in the form's order (``get_case_updates``, ``order_updates``): blocks of two cases in another order, and one
    case's update written before its create, are what HQ applies alike. A case's blocks are keyed by where the form
    holds each, so one case's two updates HQ applies in another order are one order difference (the later one
    wins), two rows of one repeat paired by position, and a block for one more case is that one block at its place.
    A case's own fields keep their names, a property is ``*``, and a block with an empty case id, which HQ refuses,
    is kept apart."""
    patient = _applied_block("p", create={"case_type": "patient", "case_name": "a"})
    visit = _applied_block("v", update={"note": "x"})
    visit_create = _applied_block("v", create={"case_type": "visit", "case_name": "b"})
    assert _case_classes(_processed_record(patient, visit), _processed_record(visit, patient)) == set()
    assert _case_classes(_processed_record(visit_create, visit), _processed_record(visit, visit_create)) == set()
    first, second = _applied_block("v", update={"note": "x"}), _applied_block("v", update={"note": "y"})
    places = [["__nova_operations", "first"], ["__nova_operations", "second"]]
    found = proof3.compare_case_processing(
        "d",
        "case_blocks@B",
        _processed_record(first, second, paths=places),
        _processed_record(second, first, paths=places[::-1]),
    )
    assert [(d.path, d.kind, d.before, d.after) for d in found] == [
        (
            "/runs/*/blocks/*/order()",
            "changed",
            ["__nova_operations/first/case", "__nova_operations/second/case"],
            ["__nova_operations/second/case", "__nova_operations/first/case"],
        )
    ]
    # Two rows of one repeat share HQ's path, so they pair by position.
    assert _case_classes(_processed_record(first, second), _processed_record(second, first)) == {
        ("/runs/*/blocks/*/case/*/update/*", "changed")
    }
    renamed = _applied_block("p", create={"case_type": "patient", "case_name": "  a  "})
    assert _case_classes(
        _processed_record(patient, visit), _processed_record(renamed, _applied_block("new"), visit)
    ) == {
        ("/runs/*/blocks/*/case/*/create/case_name", "changed"),
        ("/runs/*/blocks/*/case", "added"),
    }
    assert _case_classes(_processed_record(patient), _processed_record(patient, _applied_block(""))) == {
        ("/runs/*/blocks//update/case", "added")
    }
    refused = {"class": "IllegalCaseId", "message": "case_id must not be empty"}
    assert _case_classes(_processed_record(patient), _processed_record(patient, refusal=refused)) == {
        ("/runs/*/refusal/IllegalCaseId", "added")
    }


def test_one_block_more_or_less_in_a_case_is_that_block_wherever_it_sits_in_the_cases_sequence():
    """A guard's block that lost its case id leaves its case for the empty id's group: that is the one block, at
    its place, whether the case's other blocks follow it (an index, a close, another update) or not; none of them
    moves."""
    refused = {"class": "IllegalCaseId", "message": "case_id must not be empty"}
    first = _applied_block("v", update={"note": "x"})
    guard = _applied_block("v", update="")
    link = _applied_block("v", update={"case_type": "visit"}, index={"parent": {"@case_type": "patient"}})
    finish = _applied_block("v", update={"final": "y"}, close="")
    lost = dict(guard, **{"@case_id": ""})
    link_at, finish_at = ["__nova_operations", "link"], ["__nova_operations", "finish"]
    expected = {
        ("/runs/*/blocks/*/__nova_operations~1__nova_guard_*~1case", "removed"),
        (GUARD_KEY, "added"),
        ("/runs/*/refusal/IllegalCaseId/update/__nova_operations~1__nova_guard_*~1case", "added"),
    }
    for blocks, places in (
        ((first, guard, link, finish), ([], GUARD, link_at, finish_at)),
        ((first, link, finish, guard), ([], link_at, finish_at, GUARD)),
    ):
        before = _processed_record(*blocks, paths=places)
        after = _processed_record(
            *(lost if block is guard else block for block in blocks), paths=places, refusal=refused
        )
        assert _case_classes(before, after) == expected


def test_a_block_one_build_writes_where_the_other_writes_none_moved_and_is_compared_with_the_one_it_replaces():
    """HQ reads no block by where the form holds it: HQ's build writes a child case's block in ``subcase_<n>`` and
    Nova's local archive in the repeat's row, so the two are one block compared as itself (a trimmed name is that
    name, at the baseline's place); but a guard's block that left its case for the empty id, beside an operation's
    block written for the case where none was, are two blocks, since the guard's place is still written."""
    child = _applied_block("c", create={"case_type": "child", "case_name": "  proof  "})
    trimmed = _applied_block("c", create={"case_type": "child", "case_name": "proof"})
    found = proof3.compare_case_processing(
        "d",
        "case_blocks@local.ccz",
        _processed_record(child, paths=[["records", "item", "subcase_1"]]),
        _processed_record(trimmed, paths=[["records", "item"]]),
    )
    assert [(d.path, d.at) for d in found] == [
        (
            "/runs/*/blocks/*/*~1item~1subcase_*~1case/*/create/case_name",
            "/runs/0/blocks/c/records~1item~1subcase_1~1case/0/create/case_name",
        )
    ]
    refused = {"class": "IllegalCaseId", "message": "case_id must not be empty"}
    guard = _applied_block("v", update={"visit_status": "a"})
    tag = _applied_block("v", update={"tag": "t"})
    assert _case_classes(
        _processed_record(guard, paths=[GUARD]),
        _processed_record(dict(guard, **{"@case_id": ""}), tag, paths=[GUARD, OPERATION], refusal=refused),
    ) == {
        ("/runs/*/blocks/*/__nova_operations~1__nova_guard_*~1case", "removed"),
        ("/runs/*/blocks/*/__nova_operations~1*~1case", "added"),
        (GUARD_KEY, "added"),
        ("/runs/*/refusal/IllegalCaseId/update/__nova_operations~1__nova_guard_*~1case", "added"),
    }


def test_the_cases_a_refused_submission_leaves_untouched_are_that_refusals_by_its_cause():
    """HQ applies no case of a submission it refuses, so every case the other side's processing touched is a
    difference of the refusal, under its path: a guard's empty case id and an operation's are two classes, and a
    case both sides touched is compared as itself."""
    refused = {"class": "IllegalCaseId", "message": "case_id must not be empty"}
    patient = _applied_block("p", update={"code": "b"})
    touched = {"p": {"type": "patient", "properties": {"code": "b"}}}
    guard = _applied_block("", update={"visit_status": "a"})
    promote = _applied_block("", create={"case_type": "visit", "case_name": "n"})
    for lost, place, written in (
        (guard, GUARD, "update/__nova_operations~1__nova_guard_*~1case"),
        (promote, OPERATION, "create/__nova_operations~1*~1case"),
    ):
        found = proof3.compare_case_processing(
            "d",
            "case_blocks@B",
            _processed_record(patient, cases=touched),
            _processed_record(patient, lost, paths=[[], place], refusal=refused),
        )
        refusal = f"/runs/*/refusal/IllegalCaseId/{written}"
        assert {(d.path, d.kind) for d in found} == {
            (f"/runs/*/blocks//{written}", "added"),
            (refusal, "added"),
            (f"{refusal}/cases/*", "removed"),
        }
        action = written.split("/", 1)[0]
        assert [d.at for d in found if d.path == f"{refusal}/cases/*"] == [
            f"/runs/0/refusal/IllegalCaseId/{action}/{'~1'.join([*place, 'case'])}/cases/p"
        ]
    changed = {"p": {"type": "patient", "properties": {"code": "c"}}}
    assert _case_classes(_processed_record(patient, cases=touched), _processed_record(patient, cases=changed)) == {
        ("/runs/*/cases/*/properties/*", "changed")
    }


# Where Nova writes a guard's block and an operation's (``lib/commcare/xform/caseOps.ts``), as HQ's path to each.
GUARD = ["__nova_operations", "__nova_guard_1270f255_e852_5981_a7fd_9b23c7d98225_text"]
OPERATION = ["__nova_operations", "promote"]
# A guard's block writes an existing case and an operation's here opens one: the empty id's blocks are by the
# action HQ applies each at first.
GUARD_KEY = "/runs/*/blocks//update/__nova_operations~1__nova_guard_*~1case"
OPERATION_KEY = "/runs/*/blocks//create/__nova_operations~1*~1case"
CREATED_GUARD_KEY = "/runs/*/blocks//create/__nova_operations~1__nova_guard_*~1case"


def test_blocks_with_an_empty_case_id_are_keyed_by_where_the_form_holds_them_and_so_is_hqs_refusal():
    """HQ refuses a block whose case id is empty (``AbstractCaseDbCache.get``) and applies the empty id before
    every other (``get_case_updates`` sorts by case id), so it refuses the first empty-id block in its order. The
    empty id's blocks are keyed by the action HQ applies each at first and where the form holds each, and the
    refusal by that block, so a guard that lost its case id and an operation that did are two classes, and so are
    a block that opens a case and one that writes an existing case; inside such a block a case's own fields keep
    their names and a property is ``*``, as in any block."""
    refused = {"class": "IllegalCaseId", "message": "case_id must not be empty"}
    patient = _applied_block("p", update={"code": "b"})
    guard = _applied_block("", update={"visit_status": "a"})
    assert _case_classes(
        _processed_record(patient), _processed_record(patient, guard, paths=[[], GUARD], refusal=refused)
    ) == {
        (GUARD_KEY, "added"),
        ("/runs/*/refusal/IllegalCaseId/update/__nova_operations~1__nova_guard_*~1case", "added"),
    }
    # An operation's block with a create lost its id as well: HQ applies a create before an update, so it refuses
    # the operation's block, and the guard's is still its own class.
    promote = _applied_block("", create={"case_type": "visit", "case_name": "n"})
    assert _case_classes(
        _processed_record(patient),
        _processed_record(patient, guard, promote, paths=[[], GUARD, OPERATION], refusal=refused),
    ) == {
        (GUARD_KEY, "added"),
        (OPERATION_KEY, "added"),
        ("/runs/*/refusal/IllegalCaseId/create/__nova_operations~1*~1case", "added"),
    }
    # Both sides hold the same empty-id guard, its case's name, a property and an update different.
    before = _applied_block("", create={"case_name": "a", "visit_owner": "x"}, update={"visit_status": "a"})
    after = _applied_block("", create={"case_name": "b", "visit_owner": "y"}, update={"visit_status": "b"})
    assert _case_classes(
        _processed_record(before, paths=[GUARD], refusal=refused),
        _processed_record(after, paths=[GUARD], refusal=refused),
    ) == {
        (f"{CREATED_GUARD_KEY}/*/create/case_name", "changed"),
        (f"{CREATED_GUARD_KEY}/*/create/*", "changed"),
        (f"{CREATED_GUARD_KEY}/*/update/*", "changed"),
    }
    # HQ refuses another block on the other side: the refusal's cause moved.
    assert _case_classes(
        _processed_record(guard, paths=[GUARD], refusal=refused),
        _processed_record(guard, promote, paths=[GUARD, OPERATION], refusal=refused),
    ) == {
        (OPERATION_KEY, "added"),
        ("/runs/*/refusal/IllegalCaseId/update/__nova_operations~1__nova_guard_*~1case", "removed"),
        ("/runs/*/refusal/IllegalCaseId/create/__nova_operations~1*~1case", "added"),
    }


def test_a_record_of_hqs_blocks_without_where_the_form_holds_them_is_incomplete():
    """The observation records where the form holds each block beside the blocks; a record without them cannot
    key an empty case id's blocks, so the judge says what is missing rather than keying them all as one."""
    held = _processed_record(_applied_block("", update={"visit_status": "a"}), paths=[GUARD])
    bare = copy.deepcopy(held)
    del bare["runs"][0]["blockPaths"]
    with pytest.raises(proof3.RecordIncomplete, match="blockPaths"):
        proof3.compare_case_processing("d", "case_blocks@B", held, bare)


def test_hqs_processing_names_generated_ids_as_the_runs_trace_does():
    """A case HQ made from a generated id is keyed by the name the trace gives the id, so the same case made from
    ids the runner numbered in another order is the same key on both sides."""

    def trace(child, instance):
        submission = (
            f'<data xmlns="f"><case xmlns="{CASE_XMLNS}" case_id="@generated:uuid:{child}"/>'
            f'<meta xmlns="http://openrosa.org/jr/xforms"><instanceID>@generated:uuid:{instance}</instanceID></meta></data>'
        )
        return {"runs": [{"end": "submitted", "trace": [{"screen": "form", "submission": submission}]}]}

    def processed(child):
        mark = f"@generated:uuid:{child}"
        return _processed_record(_applied_block(mark, create={"case_type": "t"}), cases={mark: {"type": "t"}})

    assert _case_classes(processed(1), processed(2), traces=(trace(1, 2), trace(2, 1))) == set()
    assert _case_classes(processed(1), processed(2)) == {
        ("/runs/*/blocks/*/case", "added"),
        ("/runs/*/blocks/*/case", "removed"),
        ("/runs/*/cases/*", "added"),
        ("/runs/*/cases/*", "removed"),
    }


def test_a_refusal_the_bar_reports_is_not_reported_again_and_any_other_names_its_cause():
    """Where A builds nothing a runtime installs, the bar holds why (``validate_app@A``) and proof 3 adds nothing;
    the causes the bar does not hold are each a class of their own: B's errors at the bar's own path, what a step
    raised, and an archive the document does not carry."""
    from proof.checks import bar
    from proof.checks.observations import AppState, Republish
    from proof.observe.outcome import BuildOutcome

    error = {"type": "blank form", "module": {"id": 0}, "form": {"id": 1}, "message": "blank"}
    unbuilt = BuildOutcome("A", 1, [error], {}, {"suite.xml": b"<suite/>"})
    observed = Republish("d", "minimum", a=AppState("A", "app", {}, {}, unbuilt))
    assert proof3.behavioral_equivalence("d", observed, None, Blobs(), has_local=True) == []
    assert [(d.artifact, d.path) for d in bar.build_differences("d", unbuilt)] == [
        ("validate_app@A", "/modules/*/forms/*/blank form")
    ]

    raised = {"class": "XFormValidationError", "message": "no"}
    unbuilt_b = BuildOutcome("B-aligned", 1, [error], {"create_all_files": raised}, None)
    assert [(d.path, d.at, d.after) for d in proof3.unbuildable_refusals("d", "trace@B", unbuilt_b)] == [
        (
            "/unbuildable/validate_app/modules/*/forms/*/blank form",
            "/unbuildable/validate_app/modules/0/forms/1/blank form",
            error,
        ),
        (
            "/unbuildable/raised/create_all_files/XFormValidationError",
            "/unbuildable/raised/create_all_files/XFormValidationError",
            raised,
        ),
    ]

    built = BuildOutcome("A", 1, [], {}, {"suite.xml": b"<suite/>"})
    observed = Republish("d", "minimum", a=AppState("A", "app", {}, {}, built))
    assert [(d.path, d.kind) for d in proof3.behavioral_equivalence("d", observed, None, Blobs(), has_local=False)] == [
        ("/no-local-archive", "refused")
    ]
